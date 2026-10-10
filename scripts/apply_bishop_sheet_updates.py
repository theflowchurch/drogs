"""Apply personal details from the user's per-bishop pastor sheets to existing bishop records.

Only the bishop's own row (the first data row of each sheet) is used. Name, role,
organization and denomination are deliberately left unchanged; blanks and
placeholders such as MISSING or N/A never overwrite existing values.

Usage: python scripts/apply_bishop_sheet_updates.py /path/to/sheets/folder
"""
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = 'Bishop pastor sheet supplied by user, 10 October 2026'
# Sheet file name fragment -> (bishop code, current directory name).
SHEETS = {
    'Bishop_David_Yalleh': (35, 'David Yalleh'),
    'Bishop_Ebo_Ankrah': (329, 'Ebo Ankrah'),
    'Bishop_Toss_Mills-Odoi': (321, 'Toss Mills-Odoi'),
    'Bishop_Emmanuel_Opoku_Baah': (211, 'Emmanuel Opoku Baah'),
    'Bishop_Samuel_Gaspar': (177, 'Samuel Kwame Gasper'),
    'Bishop_Kent_Njeru': (222, 'Kent Njeru'),
    'Mother_Vivaldi-Joy_Aboagye': (187, 'Vivaldi-Joy Aboagye'),
    'Mother_Daniella_Hart': (208, 'Daniella Hart'),
    'Bishop_Abraham_Aryee': (2, 'Abraham Nii Ayi Aryee'),
    'Episcopal_Sister_Adelaide_Heward-Mills': (4, 'Adelaide Heward-Mills'),
    'Bishop_Stephen_Jaiden_Asamoah': (241, 'Stephen Jaiden Asamoah'),
    'Bishop_Kwabena_Amoako-Doku': (118, 'Kwabena Oware Amoako-Doku'),
    'Bishop_Roger_Hiltbrunner': (235, 'Roger Hiltbrunner'),
}
PLACEHOLDERS = {'', 'missing', 'n/a', 'na', 'email', '-', '0', 'none'}
COUNTRIES = {'drc': 'Congo, Democratic Republic', 'ivory coast': "Cote D'ivoire (ivory Coast)"}
COLUMNS = ['name', 'gender', 'rank', 'mobile', 'dateOfBirth', 'age', 'denomination', 'branch', 'city',
           'region', 'yearAppointed', 'yearOrdained', 'yearConsecrated', 'photo', 'email']
UPDATED = ['gender', 'mobile', 'dateOfBirth', 'age', 'branch', 'city', 'region',
           'yearAppointed', 'yearOrdained', 'yearConsecrated', 'email']


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    value = re.sub(r'\s+', ' ', str(value)).strip()
    return '' if value.lower() in PLACEHOLDERS else value


def birth_date(value):
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, (int, float)) and value > 10000:
        return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(value))).isoformat()
    value = text(value)
    match = re.fullmatch(r'(\d{1,2})/(\d{1,2})/(\d{4})', value)
    return datetime.date(int(match[3]), int(match[2]), int(match[1])).isoformat() if match else value


def clean(field, value):
    if field == 'dateOfBirth':
        return birth_date(value)
    value = text(value)
    if field == 'gender':
        return value.upper()
    if field == 'region':
        return COUNTRIES.get(value.lower(), value.title() if value.isupper() else value)
    if field == 'email':
        return value.lower()
    if field == 'city' and value.isupper():
        return value.title()
    return re.sub(r'\s*,\s*', ', ', value) if field == 'branch' else value


def bishop_row(path):
    sheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    header = next(i for i, row in enumerate(rows) if text(row[0]).lower() == 'full name')
    return {field: clean(field, value) for field, value in zip(COLUMNS, rows[header + 1])}


def main(folder):
    path = ROOT / 'data/bishops.js'
    bishops = json.loads(path.read_text().split('=', 1)[1].strip().rstrip(';'))
    by_code = {b['code']: b for b in bishops}
    for fragment, (code, name) in SHEETS.items():
        files = sorted(Path(folder).glob(f'*{fragment}*.xlsx'))
        if len(files) != 1:
            raise ValueError(f'Expected one sheet for {fragment}, found {len(files)}')
        person = by_code[code]
        if person['name'] != name:
            raise ValueError(f'Name changed for bishop {code}; review before applying')
        row = bishop_row(files[0])
        changes = {f: row[f] for f in UPDATED if row[f] and row[f] != person.get(f, '')}
        for field, value in changes.items():
            print(f'{code} {name}: {field} {person.get(field, "")!r} -> {value!r}')
        person.update(changes, profileSource=SOURCE)
    path.write_text('window.BISHOPS=' + json.dumps(bishops, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(f'Updated personal details for {len(SHEETS)} bishops; names, roles and denominations preserved.')


if __name__ == '__main__':
    main(sys.argv[1])
